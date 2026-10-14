"""Ingestion (PDF/DOCX/TXT/MD/images) and retrieval."""
import io

import numpy as np

import config as cfg
import db
import llm


def _chunk(text: str, size=cfg.CHUNK_SIZE, overlap=cfg.CHUNK_OVERLAP):
    text = " ".join(text.split())
    out, i = [], 0
    while i < len(text):
        end = min(i + size, len(text))
        if end < len(text):  # prefer to break at sentence end
            cut = text.rfind(". ", i + size // 2, end)
            end = cut + 1 if cut != -1 else end
        piece = text[i:end].strip()
        if piece:
            out.append(piece)
        if end >= len(text):
            break
        i = max(end - overlap, i + 1)
    return out


def _extract(name: str, data: bytes, mime: str):
    n = name.lower()
    if mime.startswith("image/"):
        return "image", [(None, llm.describe_image(data, mime))]
    if n.endswith(".pdf"):
        from pypdf import PdfReader
        return "pdf", [(i + 1, p.extract_text() or "") for i, p in enumerate(PdfReader(io.BytesIO(data)).pages)]
    if n.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                text += "\n" + " | ".join(c.text for c in row.cells)
        return "docx", [(None, text)]
    if n.endswith((".txt", ".md", ".csv")):
        return "text", [(None, data.decode("utf-8", errors="ignore"))]
    raise ValueError("unsupported file type")


def ingest(uid: int, name: str, data: bytes, mime: str) -> int:
    kind, pages = _extract(name, data, mime)
    parts = [(pg, c) for pg, txt in pages for c in _chunk(txt)]
    if not parts:
        raise ValueError("no extractable text (scanned PDF? upload pages as images instead)")
    embs = llm.embed([t for _, t in parts], "RETRIEVAL_DOCUMENT")
    db.add_document(uid, name, kind, parts, embs)
    return len(parts)


def retrieve(uid: int, query: str, k=cfg.TOP_K, min_score=cfg.MIN_SCORE):
    rows, mat = db.load_chunks(uid)
    if not rows:
        return []
    q = llm.embed([query], "RETRIEVAL_QUERY")[0]
    scores = mat @ q
    return [dict(rows[i], score=float(scores[i])) for i in np.argsort(-scores)[:k] if scores[i] >= min_score]
